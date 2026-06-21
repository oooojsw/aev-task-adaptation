from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


TRANSIENT_ERROR_CATEGORIES = {"rate_limit", "timeout", "network", "api_error"}


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig", errors="replace")


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(read_text(path))


def normalize_path(value: str | None) -> str:
    if not value:
        return ""
    return str(value).replace("\\", "/")


def parse_validator_stderr(stderr: str) -> dict[str, str]:
    out: dict[str, str] = {}
    patterns = {
        "expectedAnswer": r"Expected:\s*(.+)",
        "predictedAnswer": r"Predicted:\s*(.+)",
        "tolerance": r"Tolerance:\s*(.+)",
        "validatorScoreText": r"Score:\s*(.+)",
    }
    for key, pattern in patterns.items():
        match = re.search(pattern, stderr or "", re.IGNORECASE)
        if match:
            out[key] = match.group(1).strip()
    return out


def load_private_answers(task_pack_root: Path) -> dict[str, str]:
    answers: dict[str, str] = {}
    root = task_pack_root / "private-answers"
    if not root.exists():
        return answers
    for case_dir in sorted(root.iterdir()):
        answer_path = case_dir / "answer.txt"
        if case_dir.is_dir() and answer_path.exists():
            answers[case_dir.name] = read_text(answer_path).strip()
    return answers


def load_case_metadata(task_pack_root: Path) -> dict[str, dict[str, str]]:
    metadata: dict[str, dict[str, str]] = {}
    tasks_root = task_pack_root / "tasks"
    if not tasks_root.exists():
        return metadata
    for task_dir in sorted(tasks_root.iterdir()):
        if not task_dir.is_dir():
            continue
        case_id = task_dir.name
        item: dict[str, str] = {"caseID": case_id}
        spec_path = task_dir / "spec.json"
        if spec_path.exists():
            try:
                spec = read_json(spec_path)
                item["title"] = str(spec.get("title", ""))
                item["goal"] = str(spec.get("goal", ""))
                item["taskFamily"] = str(spec.get("taskFamily", ""))
                materials = spec.get("materials", [])
                if isinstance(materials, list):
                    item["materialCount"] = str(len(materials))
                    item["materials"] = "\n".join(
                        str(m.get("path", "")) for m in materials if isinstance(m, dict)
                    )
            except Exception as exc:
                item["metadataError"] = str(exc)
        instruction_path = task_dir / "instruction.md"
        if instruction_path.exists():
            instruction = read_text(instruction_path).strip()
            item["question"] = instruction
        metadata[case_id] = item
    return metadata


def load_prompt_audit(case_dir: Path) -> dict[str, Any]:
    for name in ("prompt-audit.json", "prompt.audit.json"):
        path = case_dir / name
        if path.exists():
            try:
                return read_json(path)
            except Exception:
                return {}
    return {}


def load_material_audit(run_dir: Path, case_id: str) -> dict[str, Any]:
    root = run_dir / "material-audit"
    if not root.exists():
        return {}
    candidates = [
        root / f"{case_id}.json",
        root / f"{case_id}.material-audit.json",
    ]
    candidates.extend(root.glob(f"*{case_id}*.json"))
    for path in candidates:
        if path.exists():
            try:
                return read_json(path)
            except Exception:
                return {}
    return {}


def load_run_results(run_dir: Path, source_label: str = "") -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    results_root = run_dir / "results"
    if not results_root.exists():
        return rows
    for result_path in sorted(results_root.rglob("result.json")):
        try:
            result = read_json(result_path)
        except Exception as exc:
            rows.append(
                {
                    "caseID": result_path.parent.name,
                    "modelID": "",
                    "passed": False,
                    "errorCategory": "result_json_read_error",
                    "error": str(exc),
                    "sourceRun": source_label or str(run_dir),
                    "resultPath": str(result_path),
                }
            )
            continue

        case_dir = result_path.parent
        case_id = str(result.get("caseID") or case_dir.name)
        validator_path = Path(str(result.get("validatorResultPath") or case_dir / "validator-result.json"))
        validator: dict[str, Any] = {}
        if validator_path.exists():
            try:
                validator = read_json(validator_path)
            except Exception:
                validator = {}
        stderr = str(validator.get("stderr") or result.get("error") or "")
        parsed_validator = parse_validator_stderr(stderr)
        usage = result.get("usage") if isinstance(result.get("usage"), dict) else {}
        prompt_audit = load_prompt_audit(case_dir)
        material_audit = load_material_audit(run_dir, case_id)

        row: dict[str, Any] = {
            "caseID": case_id,
            "modelID": result.get("modelID", ""),
            "provider": result.get("provider", ""),
            "passed": bool(result.get("passed")),
            "score": result.get("score", ""),
            "answer": str(result.get("answer", "")),
            "expectedAnswer": parsed_validator.get("expectedAnswer", ""),
            "predictedAnswer": parsed_validator.get("predictedAnswer", ""),
            "tolerance": parsed_validator.get("tolerance", ""),
            "validatorScoreText": parsed_validator.get("validatorScoreText", ""),
            "errorCategory": result.get("errorCategory") or ("none" if result.get("passed") else "unknown"),
            "error": str(result.get("error", "")),
            "attempts": result.get("attempts", ""),
            "latencyMs": result.get("latencyMs", ""),
            "inputTokens": usage.get("inputTokens", ""),
            "outputTokens": usage.get("outputTokens", ""),
            "totalTokens": usage.get("totalTokens", ""),
            "skillMode": result.get("skillMode", ""),
            "sourceRun": source_label or str(run_dir),
            "correctedByRetry": False,
            "resultPath": str(result_path),
            "responsePath": str(result.get("rawResponsePath", "")),
            "validatorPath": str(result.get("validatorResultPath", validator_path)),
            "promptTokenEstimate": prompt_audit.get("estimatedPromptTokens", ""),
            "materialAuditPath": "",
            "materialCount": "",
            "materialWarnings": "",
        }

        if material_audit:
            row["materialAuditPath"] = str((run_dir / "material-audit").resolve())
            materials = material_audit.get("materials") or material_audit.get("includedMaterials") or []
            if isinstance(materials, list):
                row["materialCount"] = len(materials)
            warnings = material_audit.get("warnings") or material_audit.get("errors") or []
            if isinstance(warnings, list):
                row["materialWarnings"] = "\n".join(str(x) for x in warnings[:8])
        rows.append(row)
    return rows


def load_aggregate_rows(aggregate_dir: Path) -> list[dict[str, Any]]:
    summary_path = aggregate_dir / "summary.json"
    source_details: dict[str, dict[str, Any]] = {}
    if summary_path.exists():
        summary = read_json(summary_path)
        for source_run in summary.get("sourceRuns", []) or []:
            source_run_path = Path(str(source_run))
            for row in load_run_results(source_run_path, source_label=source_run_path.name):
                source_details[str(row.get("caseID", ""))] = row

    rows: list[dict[str, Any]] = []
    csv_path = aggregate_dir / "results.csv"
    if not csv_path.exists():
        return rows
    with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
        for raw in csv.DictReader(handle):
            case_id = raw.get("caseID", "")
            base = dict(source_details.get(case_id, {}))
            base.update(
                {
                    "caseID": case_id,
                    "modelID": raw.get("modelID", base.get("modelID", "")),
                    "passed": str(raw.get("passed", "")).lower() == "true",
                    "errorCategory": raw.get("errorCategory", base.get("errorCategory", "")),
                    "answer": raw.get("answer", base.get("answer", "")),
                    "sourceRun": raw.get("sourceRun", base.get("sourceRun", aggregate_dir.name)),
                    "aggregateDir": str(aggregate_dir),
                }
            )
            rows.append(base)
    return rows


def apply_retry_corrections(
    main_rows: list[dict[str, Any]], retry_rows: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    retry_by_case = {str(row.get("caseID", "")): row for row in retry_rows}
    corrected: list[dict[str, Any]] = []
    for row in main_rows:
        category = str(row.get("errorCategory", ""))
        case_id = str(row.get("caseID", ""))
        if category in TRANSIENT_ERROR_CATEGORIES and case_id in retry_by_case:
            replacement = dict(retry_by_case[case_id])
            replacement["correctedByRetry"] = True
            replacement["originalErrorCategory"] = category
            replacement["originalSourceRun"] = row.get("sourceRun", "")
            corrected.append(replacement)
        else:
            corrected.append(row)
    return corrected


def enrich_rows(
    rows: list[dict[str, Any]],
    answers: dict[str, str],
    metadata: dict[str, dict[str, str]],
) -> list[dict[str, Any]]:
    enriched = []
    for raw in rows:
        row = dict(raw)
        case_id = str(row.get("caseID", ""))
        row["expectedAnswer"] = row.get("expectedAnswer") or answers.get(case_id, "")
        meta = metadata.get(case_id, {})
        row["question"] = meta.get("question", "")
        row["taskTitle"] = meta.get("title", "")
        row["taskFamily"] = meta.get("taskFamily", "")
        if not row.get("materialCount"):
            row["materialCount"] = meta.get("materialCount", "")
        row["materials"] = meta.get("materials", "")
        enriched.append(row)
    return enriched


def summarize_model(rows: list[dict[str, Any]]) -> dict[str, Any]:
    total = len(rows)
    passed = sum(1 for row in rows if row.get("passed"))
    failed = total - passed
    categories = Counter(str(row.get("errorCategory") or ("none" if row.get("passed") else "unknown")) for row in rows)

    def avg_number(key: str) -> float | str:
        values: list[float] = []
        for row in rows:
            value = row.get(key)
            if isinstance(value, (int, float)):
                values.append(float(value))
            elif isinstance(value, str) and value.strip():
                try:
                    values.append(float(value))
                except ValueError:
                    pass
        return round(sum(values) / len(values), 2) if values else ""

    return {
        "modelID": rows[0].get("modelID", "") if rows else "",
        "provider": rows[0].get("provider", "") if rows else "",
        "total": total,
        "passed": passed,
        "failed": failed,
        "passRate": passed / total if total else 0,
        "validationFailures": categories.get("validation", 0),
        "answerExtractionFailures": categories.get("answer_extraction", 0),
        "rateLimitFailures": categories.get("rate_limit", 0),
        "timeoutFailures": categories.get("timeout", 0),
        "avgLatencyMs": avg_number("latencyMs"),
        "avgInputTokens": avg_number("inputTokens"),
        "avgOutputTokens": avg_number("outputTokens"),
        "avgTotalTokens": avg_number("totalTokens"),
        "avgAttempts": avg_number("attempts"),
        "errorCategories": ", ".join(f"{k}:{v}" for k, v in sorted(categories.items())),
    }


def add_sheet(wb: Workbook, title: str, headers: list[str], rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx, header in enumerate(headers, start=1):
        max_len = len(header)
        for cell in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, values_only=True):
            for value in cell:
                if value is not None:
                    max_len = max(max_len, min(80, len(str(value))))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 60)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_case_comparison(model_rows: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    by_case: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    model_names = list(model_rows.keys())
    for model_name, rows in model_rows.items():
        for row in rows:
            by_case[str(row.get("caseID", ""))][model_name] = row
    comparison: list[dict[str, Any]] = []
    for case_id in sorted(by_case):
        item: dict[str, Any] = {"caseID": case_id}
        answers = []
        passed_flags = []
        for model_name in model_names:
            row = by_case[case_id].get(model_name, {})
            short = model_name[:28]
            item[f"{short} passed"] = row.get("passed", "")
            item[f"{short} answer"] = row.get("answer", "")
            item[f"{short} errorCategory"] = row.get("errorCategory", "")
            item[f"{short} latencyMs"] = row.get("latencyMs", "")
            if row:
                answers.append(str(row.get("answer", "")))
                passed_flags.append(bool(row.get("passed")))
        non_empty_answers = [answer for answer in answers if answer]
        item["modelsAgreeOnAnswer"] = len(set(non_empty_answers)) == 1 if len(non_empty_answers) > 1 else ""
        item["anyPassed"] = any(passed_flags)
        item["allPassed"] = all(passed_flags) if passed_flags else False
        item["allFailed"] = (not any(passed_flags)) if passed_flags else False
        comparison.append(item)
    return comparison


def write_report(
    out_path: Path,
    task_pack_root: Path,
    a3b_aggregate_dir: Path,
    qwen_main_dir: Path,
    qwen_retry_dir: Path,
    qwen_no_skill_main_dir: Path | None = None,
    qwen_no_skill_retry_dir: Path | None = None,
) -> None:
    answers = load_private_answers(task_pack_root)
    metadata = load_case_metadata(task_pack_root)

    a3b_rows = enrich_rows(load_aggregate_rows(a3b_aggregate_dir), answers, metadata)
    qwen_main_rows = load_run_results(qwen_main_dir, source_label=qwen_main_dir.name)
    qwen_retry_rows = load_run_results(qwen_retry_dir, source_label=qwen_retry_dir.name)
    qwen_rows = enrich_rows(apply_retry_corrections(qwen_main_rows, qwen_retry_rows), answers, metadata)

    model_rows = {
        "siliconflow-qwen3-6-35b-a3b": a3b_rows,
        "siliconflow-qwen3-5-4b-corrected": qwen_rows,
    }
    if qwen_no_skill_main_dir and qwen_no_skill_main_dir.exists():
        no_skill_main = load_run_results(qwen_no_skill_main_dir, source_label=qwen_no_skill_main_dir.name)
        no_skill_retry = (
            load_run_results(qwen_no_skill_retry_dir, source_label=qwen_no_skill_retry_dir.name)
            if qwen_no_skill_retry_dir and qwen_no_skill_retry_dir.exists()
            else []
        )
        model_rows["siliconflow-qwen3-5-4b-no-skill-corrected"] = enrich_rows(
            apply_retry_corrections(no_skill_main, no_skill_retry), answers, metadata
        )

    wb = Workbook()
    wb.remove(wb.active)

    overview_rows = []
    for label, rows in model_rows.items():
        summary = summarize_model(rows)
        summary["reportModelLabel"] = label
        summary["runSources"] = "\n".join(sorted({str(row.get("sourceRun", "")) for row in rows if row.get("sourceRun")}))
        overview_rows.append(summary)
    overview_headers = [
        "reportModelLabel",
        "modelID",
        "provider",
        "total",
        "passed",
        "failed",
        "passRate",
        "validationFailures",
        "answerExtractionFailures",
        "rateLimitFailures",
        "timeoutFailures",
        "avgLatencyMs",
        "avgInputTokens",
        "avgOutputTokens",
        "avgTotalTokens",
        "avgAttempts",
        "errorCategories",
        "runSources",
    ]
    add_sheet(wb, "Overview", overview_headers, overview_rows)

    detail_headers = [
        "modelLabel",
        "caseID",
        "passed",
        "score",
        "errorCategory",
        "answer",
        "expectedAnswer",
        "predictedAnswer",
        "tolerance",
        "attempts",
        "latencyMs",
        "inputTokens",
        "outputTokens",
        "totalTokens",
        "promptTokenEstimate",
        "provider",
        "modelID",
        "skillMode",
        "correctedByRetry",
        "originalErrorCategory",
        "sourceRun",
        "question",
        "materialCount",
        "materials",
        "responsePath",
        "validatorPath",
        "resultPath",
    ]
    detail_rows: list[dict[str, Any]] = []
    for label, rows in model_rows.items():
        for row in rows:
            item = dict(row)
            item["modelLabel"] = label
            detail_rows.append(item)
    add_sheet(wb, "Case Results", detail_headers, detail_rows)

    comparison_rows = build_case_comparison(model_rows)
    comparison_headers = list(comparison_rows[0].keys()) if comparison_rows else ["caseID"]
    add_sheet(wb, "Case Comparison", comparison_headers, comparison_rows)

    error_rows: list[dict[str, Any]] = []
    for label, rows in model_rows.items():
        counts = Counter(str(row.get("errorCategory") or ("none" if row.get("passed") else "unknown")) for row in rows)
        for category, count in sorted(counts.items()):
            error_rows.append({"modelLabel": label, "errorCategory": category, "count": count, "section": "category-count"})
        for row in rows:
            if not row.get("passed"):
                error_rows.append(
                    {
                        "modelLabel": label,
                        "caseID": row.get("caseID", ""),
                        "errorCategory": row.get("errorCategory", ""),
                        "answer": row.get("answer", ""),
                        "expectedAnswer": row.get("expectedAnswer", ""),
                        "predictedAnswer": row.get("predictedAnswer", ""),
                        "errorExcerpt": str(row.get("error", ""))[:500],
                        "section": "failed-case",
                    }
                )
    add_sheet(
        wb,
        "Errors",
        ["section", "modelLabel", "errorCategory", "count", "caseID", "answer", "expectedAnswer", "predictedAnswer", "errorExcerpt"],
        error_rows,
    )

    token_rows: list[dict[str, Any]] = []
    for label, rows in model_rows.items():
        summary = summarize_model(rows)
        token_rows.append({"section": "model-summary", "modelLabel": label, **summary})
        for row in rows:
            token_rows.append(
                {
                    "section": "case",
                    "modelLabel": label,
                    "caseID": row.get("caseID", ""),
                    "passed": row.get("passed", ""),
                    "latencyMs": row.get("latencyMs", ""),
                    "inputTokens": row.get("inputTokens", ""),
                    "outputTokens": row.get("outputTokens", ""),
                    "totalTokens": row.get("totalTokens", ""),
                    "attempts": row.get("attempts", ""),
                    "promptTokenEstimate": row.get("promptTokenEstimate", ""),
                }
            )
    add_sheet(
        wb,
        "Token Latency",
        [
            "section",
            "modelLabel",
            "caseID",
            "passed",
            "latencyMs",
            "inputTokens",
            "outputTokens",
            "totalTokens",
            "attempts",
            "promptTokenEstimate",
            "total",
            "passed",
            "failed",
            "passRate",
            "avgLatencyMs",
            "avgInputTokens",
            "avgOutputTokens",
            "avgTotalTokens",
            "avgAttempts",
        ],
        token_rows,
    )

    source_rows = [
        {
            "key": "generatedAt",
            "value": datetime.now().isoformat(timespec="seconds"),
            "note": "Report generated from existing Direct Chat run artifacts.",
        },
        {"key": "taskPackRoot", "value": str(task_pack_root), "note": ""},
        {"key": "a3bAggregateDir", "value": str(a3b_aggregate_dir), "note": "Uses results.csv plus sourceRuns for details."},
        {"key": "qwenMainDir", "value": str(qwen_main_dir), "note": "Main Qwen3.5-4B 172-case run."},
        {"key": "qwenRetryDir", "value": str(qwen_retry_dir), "note": "Transient retry rows replace matching transient failures."},
        {"key": "transientCategories", "value": ", ".join(sorted(TRANSIENT_ERROR_CATEGORIES)), "note": "Only these categories are replaced by retry results."},
    ]
    add_sheet(wb, "Run Sources", ["key", "value", "note"], source_rows)

    fields_rows = [
        {"field": "Overview", "meaning": "模型级总览，含总条数、通过数、通过率、错误类别、平均耗时和平均 token。"},
        {"field": "Case Results", "meaning": "逐条 case 明细，一行是一个模型在一个 UID 上的结果。"},
        {"field": "Case Comparison", "meaning": "同一个 UID 横向比较不同模型，方便看哪些题只有某个模型做对。"},
        {"field": "Errors", "meaning": "先列错误类别计数，再列每条失败的答案、标准答案和错误摘录。"},
        {"field": "Token Latency", "meaning": "成本和速度分析；如果某个 run 没记录 token，该单元格为空。"},
        {"field": "correctedByRetry", "meaning": "表示这条结果来自瞬时错误重试替换，不是主跑原始行。"},
        {"field": "expectedAnswer", "meaning": "来自 private-answers，仅用于离线报告和验收，不应注入模型提示词。"},
    ]
    add_sheet(wb, "Fields", ["field", "meaning"], fields_rows)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate an Excel report for Direct Chat evaluation runs.")
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument(
        "--task-pack-root",
        type=Path,
        default=Path("task-packs/real-long-tasks/officeqa-pro-learnbench"),
    )
    parser.add_argument(
        "--a3b-aggregate-dir",
        type=Path,
        default=Path("runtime-state/direct-chat-runs/20260618-a3b-aggregate-172"),
    )
    parser.add_argument(
        "--qwen-main-dir",
        type=Path,
        default=Path("runtime-state/direct-chat-runs/20260618-232802-officeqa-test172-qwen35-4b-skillopt-oracle-promptfix-full-c1"),
    )
    parser.add_argument(
        "--qwen-retry-dir",
        type=Path,
        default=Path("runtime-state/direct-chat-runs/20260619-000027-officeqa-qwen35-4b-transient-retry-c1"),
    )
    parser.add_argument(
        "--qwen-no-skill-main-dir",
        type=Path,
        default=Path("runtime-state/direct-chat-runs/20260619-officeqa-test172-qwen35-4b-no-skill-only-c1"),
    )
    parser.add_argument(
        "--qwen-no-skill-retry-dir",
        type=Path,
        default=Path("runtime-state/direct-chat-runs/20260619-officeqa-qwen35-4b-no-skill-transient-retry-c1"),
    )
    args = parser.parse_args()
    write_report(
        out_path=args.out,
        task_pack_root=args.task_pack_root,
        a3b_aggregate_dir=args.a3b_aggregate_dir,
        qwen_main_dir=args.qwen_main_dir,
        qwen_retry_dir=args.qwen_retry_dir,
        qwen_no_skill_main_dir=args.qwen_no_skill_main_dir,
        qwen_no_skill_retry_dir=args.qwen_no_skill_retry_dir,
    )
    print(str(args.out.resolve()))


if __name__ == "__main__":
    main()
