from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import re
import time
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


INFRA_CATEGORIES = {
    "api_error",
    "auth",
    "context_length",
    "network",
    "provider_5xx",
    "quota",
    "rate_limit",
    "timeout",
}
EXCEL_MAX_TEXT_LENGTH = 32_000
EXCEL_DATA_ROWS_PER_SHEET = 900_000
ILLEGAL_EXCEL_XML = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
FORMULA_PREFIXES = ("=", "+", "-", "@")


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8-sig", errors="replace"))


def try_read_json(path: Path) -> tuple[dict[str, Any], str]:
    try:
        value = read_json(path)
        if not isinstance(value, dict):
            return {}, f"Expected JSON object in {path}"
        return value, ""
    except Exception as exc:
        return {}, f"{type(exc).__name__}: {exc}"


def normalize_path(value: Any) -> str:
    return str(value or "").replace("\\", "/")


def file_hash(path_value: Any) -> str:
    try:
        path = Path(str(path_value or ""))
    except (TypeError, ValueError, OSError):
        return ""
    try:
        if not path.exists() or not path.is_file():
            return ""
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            while chunk := handle.read(1024 * 1024):
                digest.update(chunk)
        return digest.hexdigest()
    except OSError:
        return ""


def safe_int(value: Any, default: int = 0) -> int:
    try:
        return int(value)
    except (TypeError, ValueError, OverflowError):
        return default


def safe_float(value: Any, default: float = 0.0) -> float:
    try:
        number = float(value)
        return number if math.isfinite(number) else default
    except (TypeError, ValueError, OverflowError):
        return default


def safe_excel_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else ""
    if isinstance(value, (dict, list, tuple, set)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    text = ILLEGAL_EXCEL_XML.sub("", str(value))
    if len(text) > EXCEL_MAX_TEXT_LENGTH:
        text = text[: EXCEL_MAX_TEXT_LENGTH - 18] + "\n...[TRUNCATED]"
    if text.startswith(FORMULA_PREFIXES):
        return "'" + text
    return text


def classify_run(name: str, schema: str, total: int) -> str:
    lowered = name.lower()
    if "aborted" in lowered:
        return "aborted"
    if "smoke" in lowered:
        return "smoke"
    if "aggregate" in lowered or "aggregate" in schema:
        return "aggregate"
    if any(token in lowered for token in ("retry", "resume", "missing", "final2", "delta")):
        return "retry-or-delta"
    if total >= 100:
        return "full-run"
    return "subset"


def load_case_rows(run_dir: Path, summary: dict[str, Any]) -> list[dict[str, Any]]:
    if isinstance(summary.get("results"), list):
        return [dict(item) for item in summary["results"] if isinstance(item, dict)]

    rows: list[dict[str, Any]] = []
    results_root = run_dir / "results"
    if results_root.exists():
        for result_path in sorted(results_root.rglob("result.json")):
            result, error = try_read_json(result_path)
            if error:
                rows.append(
                    {
                        "caseID": result_path.parent.name,
                        "passed": False,
                        "errorCategory": "result_json_read_error",
                        "error": error,
                    }
                )
            else:
                rows.append(result)
    if rows:
        return rows

    csv_path = run_dir / "results.csv"
    if csv_path.exists():
        with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
            for raw in csv.DictReader(handle):
                item = dict(raw)
                item["passed"] = str(item.get("passed", "")).lower() == "true"
                rows.append(item)
    return rows


def case_dir_for_result(run_dir: Path, result: dict[str, Any]) -> Path | None:
    raw_response = result.get("rawResponsePath")
    if raw_response:
        try:
            path = Path(str(raw_response))
            if path.exists():
                return path.parent
        except (TypeError, ValueError, OSError):
            pass
    model_id = str(result.get("modelID", ""))
    case_id = str(result.get("caseID", ""))
    if model_id and case_id:
        candidate = run_dir / "results" / model_id / case_id
        if candidate.exists():
            return candidate
    return None


def prompt_audit_fields(run_dir: Path, result: dict[str, Any]) -> dict[str, Any]:
    case_dir = case_dir_for_result(run_dir, result)
    if not case_dir:
        return {
            "reasoningTokens": "",
            "estimatedPromptTokens": "",
            "promptBudgetTokens": "",
            "skillTokens": "",
            "materialTokens": "",
            "materialCount": "",
            "skillTruncated": "",
            "materialTruncated": "",
            "usageSource": "provider" if isinstance(result.get("usage"), dict) else "missing",
            "costUSD": "",
            "promptAuditPath": "",
        }
    audit_path = case_dir / "prompt-audit.json"
    audit, error = try_read_json(audit_path)
    usage = result.get("usage") if isinstance(result.get("usage"), dict) else {}
    if error:
        return {
            "reasoningTokens": usage.get("reasoningTokens", ""),
            "estimatedPromptTokens": "",
            "promptBudgetTokens": "",
            "skillTokens": "",
            "materialTokens": "",
            "materialCount": "",
            "skillTruncated": "",
            "materialTruncated": "",
            "usageSource": "provider" if usage else "missing",
            "costUSD": "",
            "promptAuditPath": normalize_path(audit_path),
        }
    skill = audit.get("skill") if isinstance(audit.get("skill"), dict) else {}
    materials = audit.get("materials") if isinstance(audit.get("materials"), list) else []
    material_tokens = 0
    material_truncated = 0
    for item in materials:
        if not isinstance(item, dict):
            continue
        material_tokens += safe_int(item.get("includedTokens"), 0)
        if item.get("truncated"):
            material_truncated += 1
    usage_source = "provider" if usage else "estimate" if audit.get("estimatedPromptTokens") else "missing"
    return {
        "reasoningTokens": usage.get("reasoningTokens", ""),
        "estimatedPromptTokens": audit.get("estimatedPromptTokens", ""),
        "promptBudgetTokens": audit.get("promptBudgetTokens", ""),
        "skillTokens": skill.get("includedTokens", ""),
        "materialTokens": material_tokens,
        "materialCount": len(materials),
        "skillTruncated": bool(skill.get("truncated")),
        "materialTruncated": material_truncated,
        "usageSource": usage_source,
        "costUSD": "",
        "promptAuditPath": normalize_path(audit_path),
    }


def summary_items(summary: dict[str, Any]) -> list[dict[str, Any]]:
    if isinstance(summary.get("summary"), list):
        return [dict(item) for item in summary["summary"] if isinstance(item, dict)]
    if "total" in summary:
        return [
            {
                "modelID": summary.get("modelID", ""),
                "total": summary.get("total", 0),
                "passed": summary.get("passed", 0),
                "failed": summary.get("failed", 0),
                "passRate": summary.get("passRate", summary.get("accuracy", 0)),
                "avgLatencyMs": summary.get("avgLatencyMs", ""),
            }
        ]
    return []


def scan_runs(runs_root: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    run_rows: list[dict[str, Any]] = []
    case_rows: list[dict[str, Any]] = []
    hash_cache: dict[str, str] = {}

    if not runs_root.exists():
        return [
            {
                "runID": "",
                "runType": "missing-runs-root",
                "status": "report-warning",
                "reportWarnings": f"Runs root does not exist: {normalize_path(runs_root)}",
                "runPath": normalize_path(runs_root),
            }
        ], []

    try:
        run_dirs = sorted((path for path in runs_root.iterdir() if path.is_dir()), key=lambda p: p.name)
    except OSError as exc:
        return [
            {
                "runID": "",
                "runType": "unreadable-runs-root",
                "status": "report-warning",
                "reportWarnings": f"{type(exc).__name__}: {exc}",
                "runPath": normalize_path(runs_root),
            }
        ], []

    for run_dir in run_dirs:
        summary_path = run_dir / "summary.json"
        if not summary_path.exists():
            continue
        summary, summary_error = try_read_json(summary_path)
        if summary_error:
            run_rows.append(
                {
                    "runID": run_dir.name,
                    "runType": "unreadable",
                    "status": "summary-read-error",
                    "reportWarnings": summary_error,
                    "runPath": normalize_path(run_dir),
                    "summaryPath": normalize_path(summary_path),
                }
            )
            continue

        config_path = run_dir / "config.snapshot.json"
        config: dict[str, Any] = {}
        config_error = ""
        if config_path.exists():
            config, config_error = try_read_json(config_path)
        dataset_path = normalize_path(config.get("datasetBundlePath", summary.get("datasetBundlePath", "")))
        skill_path = normalize_path(config.get("skillBundlePath", summary.get("skillBundlePath", "")))
        for value in (dataset_path, skill_path):
            if value and value not in hash_cache:
                hash_cache[value] = file_hash(value)
        try:
            completed_at = datetime.fromtimestamp(summary_path.stat().st_mtime).isoformat(timespec="seconds")
        except OSError:
            completed_at = ""
        try:
            raw_cases = load_case_rows(run_dir, summary)
        except OSError as exc:
            raw_cases = []
            config_error = "; ".join(
                value for value in (config_error, f"Case scan failed: {type(exc).__name__}: {exc}") if value
            )
        validity = summary.get("validity") if isinstance(summary.get("validity"), dict) else {}
        validity_status = str(validity.get("status") or "legacy-unverified")
        validity_reasons = validity.get("reasons") if isinstance(validity.get("reasons"), list) else []
        for row in raw_cases:
            row.update(prompt_audit_fields(run_dir, row))

        for item in summary_items(summary):
            model_id = str(item.get("modelID", ""))
            model_cases = [row for row in raw_cases if not model_id or str(row.get("modelID", "")) == model_id]
            token_input = sum(safe_int((row.get("usage") or {}).get("inputTokens") if isinstance(row.get("usage"), dict) else row.get("inputTokens"), 0) for row in model_cases)
            token_output = sum(safe_int((row.get("usage") or {}).get("outputTokens") if isinstance(row.get("usage"), dict) else row.get("outputTokens"), 0) for row in model_cases)
            token_reasoning = sum(safe_int(row.get("reasoningTokens"), 0) for row in model_cases)
            token_total = sum(safe_int((row.get("usage") or {}).get("totalTokens") if isinstance(row.get("usage"), dict) else row.get("totalTokens"), 0) for row in model_cases)
            token_estimated_prompt = sum(safe_int(row.get("estimatedPromptTokens"), 0) for row in model_cases)
            total = safe_int(item.get("total", len(model_cases)), len(model_cases))
            categories = Counter(str(row.get("errorCategory") or ("none" if row.get("passed") else "unknown")) for row in model_cases)
            skill_modes = sorted({str(row.get("skillMode", "")) for row in model_cases if row.get("skillMode")})
            skill_tiers = sorted({str(row.get("skillTier", "")) for row in model_cases if row.get("skillTier")})
            result_skill_hashes = sorted({str(row.get("skillHash", "")) for row in model_cases if row.get("skillHash")})
            providers = sorted({str(row.get("provider", "")) for row in model_cases if row.get("provider")})
            infra_count = sum(categories.get(category, 0) for category in INFRA_CATEGORIES)
            run_type = classify_run(run_dir.name, str(summary.get("schema", "")), total)
            run_rows.append(
                {
                    "runID": run_dir.name,
                    "completedAt": completed_at,
                    "runType": run_type,
                    "status": (
                        "invalid"
                        if validity_status != "valid"
                        else "complete-with-infra-failures"
                        if infra_count
                        else "complete"
                    ),
                    "validityStatus": validity_status,
                    "validityReasons": "\n".join(str(value) for value in validity_reasons),
                    "modelID": model_id,
                    "provider": ", ".join(providers),
                    "skillMode": ", ".join(skill_modes),
                    "skillTier": ", ".join(skill_tiers),
                    "total": total,
                    "passed": safe_int(item.get("passed", 0)),
                    "failed": safe_int(item.get("failed", total - safe_int(item.get("passed", 0)))),
                    "passRate": safe_float(item.get("passRate", item.get("accuracy", 0))),
                    "avgLatencyMs": item.get("avgLatencyMs", ""),
                    "inputTokens": token_input or "",
                    "outputTokens": token_output or "",
                    "reasoningTokens": token_reasoning or "",
                    "totalTokens": token_total or "",
                    "estimatedPromptTokens": token_estimated_prompt or "",
                    "costUSD": "",
                    "priceStatus": "missing",
                    "infrastructureFailures": infra_count,
                    "validationFailures": categories.get("validation", 0),
                    "parseFailures": categories.get("parse", 0) + categories.get("answer_extraction", 0),
                    "errorCategories": ", ".join(f"{key}:{value}" for key, value in sorted(categories.items())),
                    "datasetBundle": dataset_path,
                    "datasetHash": hash_cache.get(dataset_path, ""),
                    "skillBundle": skill_path,
                    "skillHash": ", ".join(result_skill_hashes) or ("NO_SKILL" if "no-skill" in skill_modes else hash_cache.get(skill_path, "")),
                    "globalConcurrency": config.get("globalConcurrency", ""),
                    "caseSelectionCount": len(config.get("caseIDs", [])) if isinstance(config.get("caseIDs"), list) else "",
                    "schema": summary.get("schema", ""),
                    "runPath": normalize_path(run_dir),
                    "summaryPath": normalize_path(summary_path),
                    "reportWarnings": config_error,
                }
            )

        for result in raw_cases:
            usage = result.get("usage") if isinstance(result.get("usage"), dict) else {}
            case_rows.append(
                {
                    "runID": run_dir.name,
                    "completedAt": completed_at,
                    "runType": classify_run(run_dir.name, str(summary.get("schema", "")), len(raw_cases)),
                    "caseID": result.get("caseID", ""),
                    "modelID": result.get("modelID", ""),
                    "provider": result.get("provider", ""),
                    "skillMode": result.get("skillMode", ""),
                    "skillTier": result.get("skillTier", ""),
                    "passed": bool(result.get("passed")),
                    "score": result.get("score", ""),
                    "answer": result.get("answer", ""),
                    "errorCategory": result.get("errorCategory", ""),
                    "attempts": result.get("attempts", ""),
                    "latencyMs": result.get("latencyMs", ""),
                    "inputTokens": usage.get("inputTokens", ""),
                    "outputTokens": usage.get("outputTokens", ""),
                    "reasoningTokens": result.get("reasoningTokens", ""),
                    "totalTokens": usage.get("totalTokens", ""),
                    "estimatedPromptTokens": result.get("estimatedPromptTokens", ""),
                    "promptBudgetTokens": result.get("promptBudgetTokens", ""),
                    "skillTokens": result.get("skillTokens", ""),
                    "materialTokens": result.get("materialTokens", ""),
                    "materialCount": result.get("materialCount", ""),
                    "skillTruncated": result.get("skillTruncated", ""),
                    "materialTruncated": result.get("materialTruncated", ""),
                    "usageSource": result.get("usageSource", ""),
                    "costUSD": result.get("costUSD", ""),
                    "error": str(result.get("error", ""))[:1000],
                    "datasetHash": hash_cache.get(dataset_path, ""),
                    "skillHash": result.get("skillHash") or ("NO_SKILL" if result.get("skillMode") == "no-skill" else hash_cache.get(skill_path, "")),
                    "responsePath": normalize_path(result.get("rawResponsePath", "")),
                    "validatorPath": normalize_path(result.get("validatorResultPath", "")),
                    "promptAuditPath": result.get("promptAuditPath", ""),
                    "runPath": normalize_path(run_dir),
                    "validityStatus": validity_status,
                }
            )
    return run_rows, case_rows


def load_composites(config_path: Path, runs_root: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if not config_path.exists():
        return [], []
    config, config_error = try_read_json(config_path)
    if config_error:
        return [
            {
                "experimentID": "CONFIG_ERROR",
                "label": "Official experiment configuration could not be read",
                "total": 0,
                "passed": 0,
                "failed": 0,
                "passRate": 0,
                "infrastructureFailures": 0,
                "isComplete": False,
                "missingSources": "",
                "sources": "",
                "note": config_error,
            }
        ], []
    summaries: list[dict[str, Any]] = []
    details: list[dict[str, Any]] = []
    for experiment in config.get("experiments", []):
        if not isinstance(experiment, dict):
            continue
        by_case: dict[str, dict[str, Any]] = {}
        source_names: list[str] = []
        missing_sources: list[str] = []
        invalid_sources: list[str] = []
        allowed_models = {str(value) for value in experiment.get("modelIDs", []) if value}
        allowed_providers = {str(value) for value in experiment.get("providers", []) if value}
        allowed_skill_modes = {str(value) for value in experiment.get("skillModes", []) if value}
        replace_policy = str(experiment.get("replacePolicy", "prefer-non-infrastructure"))
        for source_name in experiment.get("sources", []):
            source_dir = runs_root / str(source_name)
            summary_path = source_dir / "summary.json"
            if not summary_path.exists():
                missing_sources.append(str(source_name))
                continue
            source_names.append(source_dir.name)
            summary, summary_error = try_read_json(summary_path)
            if summary_error:
                missing_sources.append(f"{source_name} ({summary_error})")
                continue
            source_validity = summary.get("validity") if isinstance(summary.get("validity"), dict) else {}
            source_validity_status = str(source_validity.get("status") or "legacy-unverified")
            if source_validity_status != "valid":
                reasons = source_validity.get("reasons") if isinstance(source_validity.get("reasons"), list) else []
                suffix = f": {', '.join(str(value) for value in reasons)}" if reasons else ""
                invalid_sources.append(f"{source_name} ({source_validity_status}{suffix})")
            try:
                source_rows = load_case_rows(source_dir, summary)
            except OSError as exc:
                missing_sources.append(f"{source_name} ({type(exc).__name__}: {exc})")
                continue
            for row in source_rows:
                if allowed_models and str(row.get("modelID", "")) not in allowed_models:
                    continue
                if allowed_providers and str(row.get("provider", "")) not in allowed_providers:
                    continue
                if allowed_skill_modes and str(row.get("skillMode", "")) not in allowed_skill_modes:
                    continue
                case_id = str(row.get("caseID", ""))
                if case_id:
                    replacement = {**row, "sourceRun": source_dir.name}
                    previous = by_case.get(case_id)
                    previous_is_infra = bool(previous) and str(previous.get("errorCategory", "")) in INFRA_CATEGORIES
                    replacement_is_infra = str(replacement.get("errorCategory", "")) in INFRA_CATEGORIES
                    if (
                        replace_policy == "prefer-non-infrastructure"
                        and previous
                        and not previous_is_infra
                        and replacement_is_infra
                    ):
                        continue
                    by_case[case_id] = replacement
        total = len(by_case)
        passed = sum(1 for row in by_case.values() if row.get("passed"))
        infra = sum(1 for row in by_case.values() if str(row.get("errorCategory", "")) in INFRA_CATEGORIES)
        expected_cases = safe_int(experiment.get("expectedCases", total), total)
        summaries.append(
            {
                "experimentID": experiment.get("id", ""),
                "label": experiment.get("label", ""),
                "dataset": experiment.get("dataset", ""),
                "modelFamily": experiment.get("modelFamily", ""),
                "skillMode": experiment.get("skillMode", ""),
                "total": total,
                "passed": passed,
                "failed": total - passed,
                "passRate": passed / total if total else 0,
                "infrastructureFailures": infra,
                "isComplete": total == expected_cases and infra == 0 and not missing_sources and not invalid_sources,
                "validityStatus": "valid" if not invalid_sources else "invalid",
                "invalidSources": "\n".join(invalid_sources),
                "expectedCases": expected_cases,
                "missingSources": "\n".join(missing_sources),
                "replacementPolicy": replace_policy,
                "sources": "\n".join(source_names),
                "note": experiment.get("note", ""),
            }
        )
        for case_id, row in sorted(by_case.items()):
            details.append(
                {
                    "experimentID": experiment.get("id", ""),
                    "caseID": case_id,
                    "passed": bool(row.get("passed")),
                    "answer": row.get("answer", ""),
                    "errorCategory": row.get("errorCategory", ""),
                    "sourceRun": row.get("sourceRun", ""),
                    "provider": row.get("provider", ""),
                    "modelID": row.get("modelID", ""),
                    "skillMode": row.get("skillMode", ""),
                    "skillTier": row.get("skillTier", ""),
                }
            )
    return summaries, details


def add_table_sheet(wb: Workbook, title: str, headers: list[str], rows: list[dict[str, Any]]) -> Any:
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append([safe_excel_value(row.get(header, "")) for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_index, header in enumerate(headers, start=1):
        values = [header]
        values.extend(ws.cell(row=row_index, column=column_index).value for row_index in range(2, min(ws.max_row, 250) + 1))
        width = min(max(max((len(str(value)) for value in values if value is not None), default=10) + 2, 10), 48)
        ws.column_dimensions[get_column_letter(column_index)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for header_name in ("passRate",):
        if header_name in headers:
            column = headers.index(header_name) + 1
            for row_index in range(2, ws.max_row + 1):
                ws.cell(row=row_index, column=column).number_format = "0.00%"
    return ws


def add_partitioned_sheets(
    wb: Workbook,
    base_title: str,
    headers: list[str],
    rows: list[dict[str, Any]],
) -> list[Any]:
    if not rows:
        return [add_table_sheet(wb, base_title, headers, [])]
    sheets = []
    for offset in range(0, len(rows), EXCEL_DATA_ROWS_PER_SHEET):
        part = rows[offset : offset + EXCEL_DATA_ROWS_PER_SHEET]
        part_index = offset // EXCEL_DATA_ROWS_PER_SHEET
        title = base_title if part_index == 0 else f"{base_title[:25]} {part_index + 1:03d}"
        sheets.append(add_table_sheet(wb, title, headers, part))
    return sheets


def build_workbook(
    run_rows: list[dict[str, Any]],
    case_rows: list[dict[str, Any]],
    composite_rows: list[dict[str, Any]],
    composite_details: list[dict[str, Any]],
) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    dashboard = wb.create_sheet("Dashboard")
    dashboard.sheet_view.showGridLines = False
    dashboard["A1"] = "Direct Chat Evaluation Master Report"
    dashboard["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    dashboard["A1"].fill = PatternFill("solid", fgColor="17365D")
    dashboard.merge_cells("A1:H2")
    dashboard["A4"] = "Generated At"
    dashboard["B4"] = datetime.now().isoformat(timespec="seconds")
    dashboard["A5"] = "Completed Run Rows"
    dashboard["B5"] = len(run_rows)
    dashboard["A6"] = "Case Result Rows"
    dashboard["B6"] = len(case_rows)
    dashboard["A7"] = "Official Composite Experiments"
    dashboard["B7"] = len(composite_rows)
    dashboard["A9"] = "Official experiment summaries are curated from explicit source runs; all raw runs remain in Run Ledger."
    dashboard.merge_cells("A9:H9")
    dashboard["A9"].font = Font(italic=True, color="666666")
    dashboard.column_dimensions["A"].width = 34
    dashboard.column_dimensions["B"].width = 24

    official_headers = [
        "experimentID", "label", "dataset", "modelFamily", "skillMode", "total", "passed", "failed",
        "passRate", "infrastructureFailures", "isComplete", "validityStatus", "expectedCases", "missingSources",
        "invalidSources",
        "replacementPolicy", "sources", "note",
    ]
    official = add_table_sheet(wb, "Official Experiments", official_headers, composite_rows)
    if composite_rows:
        chart = BarChart()
        chart.type = "col"
        chart.title = "Official Experiment Pass Rates"
        chart.y_axis.title = "Pass Rate"
        chart.y_axis.scaling.max = 1
        chart.height = 8
        chart.width = 16
        rate_col = official_headers.index("passRate") + 1
        label_col = official_headers.index("label") + 1
        chart.add_data(Reference(official, min_col=rate_col, min_row=1, max_row=official.max_row), titles_from_data=True)
        chart.set_categories(Reference(official, min_col=label_col, min_row=2, max_row=official.max_row))
        dashboard.add_chart(chart, "D4")

    run_headers = [
        "runID", "completedAt", "runType", "status", "validityStatus", "validityReasons",
        "modelID", "provider", "skillMode", "skillTier", "total", "passed",
        "failed", "passRate", "avgLatencyMs", "inputTokens", "outputTokens", "reasoningTokens", "totalTokens",
        "estimatedPromptTokens", "costUSD", "priceStatus", "infrastructureFailures", "validationFailures", "parseFailures",
        "errorCategories", "datasetBundle", "datasetHash", "skillBundle", "skillHash", "globalConcurrency",
        "caseSelectionCount", "schema", "reportWarnings", "runPath", "summaryPath",
    ]
    add_table_sheet(wb, "Run Ledger", run_headers, sorted(run_rows, key=lambda row: (str(row.get("completedAt", "")), str(row.get("runID", ""))), reverse=True))

    case_headers = [
        "runID", "completedAt", "runType", "validityStatus", "caseID", "modelID", "provider", "skillMode", "skillTier", "passed", "score",
        "answer", "errorCategory", "attempts", "latencyMs", "inputTokens", "outputTokens", "reasoningTokens", "totalTokens",
        "estimatedPromptTokens", "promptBudgetTokens", "skillTokens", "materialTokens", "materialCount", "skillTruncated",
        "materialTruncated", "usageSource", "costUSD", "error", "datasetHash", "skillHash", "responsePath", "validatorPath",
        "promptAuditPath", "runPath",
    ]
    add_partitioned_sheets(
        wb,
        "Case Results",
        case_headers,
        sorted(case_rows, key=lambda row: (str(row.get("completedAt", "")), str(row.get("caseID", ""))), reverse=True),
    )

    add_partitioned_sheets(
        wb,
        "Official Case Results",
        ["experimentID", "caseID", "passed", "answer", "errorCategory", "sourceRun", "provider", "modelID", "skillMode", "skillTier"],
        composite_details,
    )
    add_partitioned_sheets(
        wb,
        "Errors",
        ["runID", "caseID", "modelID", "provider", "skillMode", "errorCategory", "answer", "error", "runPath"],
        [row for row in case_rows if not row.get("passed")],
    )
    add_table_sheet(
        wb,
        "Fields",
        ["sheet", "meaning"],
        [
            {"sheet": "Dashboard", "meaning": "总览和正式实验成功率图。"},
            {"sheet": "Official Experiments", "meaning": "显式来源组合后的正式对比结果，补跑按来源顺序覆盖旧结果。"},
            {"sheet": "Run Ledger", "meaning": "每次完成的测试运行占一行；烟测、补跑、正式整批和旧实验全部保留。"},
            {"sheet": "Case Results", "meaning": "所有运行的逐题原始结果；不会因为后续补跑而删除历史行。"},
            {"sheet": "Official Case Results", "meaning": "正式组合实验逐题收口结果。"},
            {"sheet": "Errors", "meaning": "所有失败结果，包含模型答错、解析失败和供应商基础设施失败。"},
            {"sheet": "status", "meaning": "complete-with-infra-failures 表示运行结束但含额度、鉴权、网络或超时问题，不能直接当模型答错。"},
            {"sheet": "datasetHash / skillHash", "meaning": "用于判断测试材料和 skill 是否真的是同一版本。"},
        ],
    )
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    return wb


def acquire_lock(lock_path: Path, timeout_seconds: int = 90) -> int:
    deadline = time.time() + timeout_seconds
    while True:
        try:
            return os.open(lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        except FileExistsError:
            try:
                if time.time() - lock_path.stat().st_mtime > 15 * 60:
                    lock_path.unlink(missing_ok=True)
                    continue
            except FileNotFoundError:
                continue
            if time.time() >= deadline:
                raise TimeoutError(f"Timed out waiting for report lock: {lock_path}")
            time.sleep(1)


def promote_pending(out_path: Path) -> bool:
    pending_path = out_path.with_name(f"{out_path.stem}.pending.xlsx")
    if not pending_path.exists():
        return False
    try:
        os.replace(pending_path, out_path)
        return True
    except PermissionError:
        return False


def save_atomic(wb: Workbook, out_path: Path) -> str:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = out_path.with_name(f"{out_path.stem}.{os.getpid()}.tmp.xlsx")
    pending_path = out_path.with_name(f"{out_path.stem}.pending.xlsx")
    wb.save(temp_path)
    for attempt in range(6):
        try:
            os.replace(temp_path, out_path)
            pending_path.unlink(missing_ok=True)
            return "updated"
        except PermissionError:
            if attempt == 5:
                os.replace(temp_path, pending_path)
                return "pending-excel-close"
            time.sleep(0.5 * (2**attempt))
    return "unknown"


def write_status(status_path: Path, payload: dict[str, Any]) -> None:
    status_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = status_path.with_name(f"{status_path.name}.{os.getpid()}.tmp")
    temp_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    os.replace(temp_path, status_path)


def main() -> None:
    auto_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description="Update the canonical Direct Chat Excel report from all completed runs.")
    parser.add_argument("--runs-root", type=Path, default=auto_root / "runtime-state" / "direct-chat-runs")
    parser.add_argument("--out", type=Path, default=auto_root / "runtime-state" / "direct-chat-reports" / "direct-chat-evaluation-master.xlsx")
    parser.add_argument("--composites-config", type=Path, default=auto_root / "direct-chat-master-report.json")
    parser.add_argument("--current-run", type=Path, default=None, help="Audit-only hint; all completed runs are still scanned.")
    args = parser.parse_args()

    lock_path = args.out.with_suffix(args.out.suffix + ".lock")
    status_path = args.out.with_name("master-report-status.json")
    args.out.parent.mkdir(parents=True, exist_ok=True)
    lock_fd = acquire_lock(lock_path)
    try:
        os.write(lock_fd, f"pid={os.getpid()} started={datetime.now().isoformat()}\n".encode("utf-8"))
        promoted_pending = promote_pending(args.out.resolve())
        run_rows, case_rows = scan_runs(args.runs_root.resolve())
        composite_rows, composite_details = load_composites(args.composites_config.resolve(), args.runs_root.resolve())
        workbook = build_workbook(run_rows, case_rows, composite_rows, composite_details)
        update_status = save_atomic(workbook, args.out.resolve())
        payload = {
            "schema": "direct-chat-master-report-status/v1",
            "updatedAt": datetime.now().isoformat(timespec="seconds"),
            "status": update_status,
            "out": normalize_path(args.out.resolve()),
            "pendingPath": normalize_path(args.out.with_name(f"{args.out.stem}.pending.xlsx").resolve()),
            "promotedPendingAtStart": promoted_pending,
            "currentRun": normalize_path(args.current_run.resolve()) if args.current_run else "",
            "runRows": len(run_rows),
            "caseRows": len(case_rows),
            "officialExperiments": len(composite_rows),
        }
        write_status(status_path, payload)
        print(json.dumps(payload, ensure_ascii=False))
    except Exception as exc:
        write_status(
            status_path,
            {
                "schema": "direct-chat-master-report-status/v1",
                "updatedAt": datetime.now().isoformat(timespec="seconds"),
                "status": "error",
                "out": normalize_path(args.out.resolve()),
                "currentRun": normalize_path(args.current_run.resolve()) if args.current_run else "",
                "error": f"{type(exc).__name__}: {exc}",
            },
        )
        raise
    finally:
        os.close(lock_fd)
        lock_path.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
