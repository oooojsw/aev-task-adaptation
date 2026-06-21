import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def cell_value(value):
    if value is None:
        return ""
    return str(value)


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--max-sheets", type=int, default=8)
    parser.add_argument("--max-rows", type=int, default=30)
    parser.add_argument("--max-cols", type=int, default=20)
    args = parser.parse_args()

    path = Path(args.xlsx)
    if not path.exists():
        raise SystemExit(f"Missing XLSX: {path}")

    wb_values = load_workbook(path, data_only=True, read_only=True)
    wb_formulas = load_workbook(path, data_only=False, read_only=True)
    sheets = []
    for sheet_name in wb_values.sheetnames[: args.max_sheets]:
        ws_values = wb_values[sheet_name]
        ws_formulas = wb_formulas[sheet_name]
        rows = []
        formula_cells = []
        max_row = min(ws_values.max_row or 0, args.max_rows)
        max_col = min(ws_values.max_column or 0, args.max_cols)
        for row in range(1, max_row + 1):
            out = []
            for col in range(1, max_col + 1):
                value = ws_values.cell(row=row, column=col).value
                formula = ws_formulas.cell(row=row, column=col).value
                out.append(cell_value(value))
                if isinstance(formula, str) and formula.startswith("="):
                    formula_cells.append({"row": row, "col": col, "formula": formula})
            rows.append(out)
        sheets.append(
            {
                "name": sheet_name,
                "maxRow": ws_values.max_row,
                "maxColumn": ws_values.max_column,
                "previewRows": rows,
                "formulaCellsPreview": formula_cells[:50],
            }
        )

    print(json.dumps({"xlsx": str(path), "sheets": sheets}, ensure_ascii=False))


if __name__ == "__main__":
    main()
