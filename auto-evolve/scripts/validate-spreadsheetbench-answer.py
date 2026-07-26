#!/usr/bin/env python
import argparse
import json
import math
import os
import re
import sys
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_json", nargs="?", default=None)
    parser.add_argument("--answer", required=True)
    parser.add_argument("--candidate", required=True)
    parser.add_argument("--sheet", default=None)
    parser.add_argument("--range", dest="cell_range", default=None)
    return parser.parse_args()


def load_payload(path_text: str | None) -> dict:
    if not path_text:
        return {}
    path = os.path.abspath(path_text)
    if not os.path.exists(path):
        return {}
    try:
        with open(path, "r", encoding="utf-8-sig") as f:
            raw = json.load(f)
        return raw if isinstance(raw, dict) else {}
    except Exception:
        return {}


def load_json_file(path: str) -> Any:
    with open(path, "rb") as f:
        data = f.read()
    encodings = ["utf-8-sig", "utf-16", "utf-16-le", "utf-16-be"]
    last_error: Exception | None = None
    for encoding in encodings:
        try:
            return json.loads(data.decode(encoding))
        except Exception as exc:
            last_error = exc
    raise RuntimeError(f"candidate json could not be decoded as UTF-8/UTF-16: {last_error}")


def infer_task_dir(payload: dict) -> str | None:
    for item in payload.get("materials", []):
        if not isinstance(item, dict):
            continue
        value = item.get("path")
        role = item.get("role")
        if isinstance(value, str) and role == "human-guidance" and os.path.exists(value):
            return os.path.dirname(value)
    for item in payload.get("materials", []):
        if not isinstance(item, dict):
            continue
        value = item.get("path")
        if not isinstance(value, str) or not os.path.exists(value):
            continue
        parent = os.path.dirname(value)
        return os.path.dirname(parent) if os.path.basename(parent).lower() == "materials" else parent
    return None


def payload_output_candidates(payload: dict) -> list[str]:
    out: list[str] = []
    for item in payload.get("targetPaths", []):
        if isinstance(item, dict) and isinstance(item.get("path"), str):
            out.append(item["path"])
    for item in payload.get("expectedArtifacts", []):
        if isinstance(item, str):
            out.append(item)
    return out


def resolve_path(value: str, payload: dict, *, prefer_payload_output: bool = False) -> str:
    if os.path.isabs(value):
        return value
    if prefer_payload_output:
        wanted = os.path.basename(value).lower()
        for item in payload_output_candidates(payload):
            if os.path.basename(item).lower() == wanted and os.path.exists(item):
                return item
        for item in payload_output_candidates(payload):
            if os.path.basename(item).lower() == wanted:
                return item
    task_dir = infer_task_dir(payload)
    if task_dir:
        candidate = os.path.abspath(os.path.join(task_dir, value))
        if os.path.exists(candidate):
            return candidate
    return os.path.abspath(value)


def norm(v: Any) -> Any:
    if v is None:
        return ""
    if isinstance(v, bool):
        return v
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        if math.isfinite(v):
            return round(v, 8)
        return str(v)
    s = str(v).strip()
    if "T" in s and len(s) >= 19 and s[4] == "-" and s[7] == "-":
      s = s.replace("T", " ")
    # normalize numeric-looking strings
    try:
        fv = float(s)
        if math.isfinite(fv):
            return round(fv, 8)
    except Exception:
        pass
    return s


def clean_sheet_name(value: str | None) -> str:
    return str(value or "").strip().strip("'").strip('"')


def clean_cell_ref(value: str | None) -> str:
    return str(value or "").strip().replace("\uff1a", ":").strip("'").strip('"')


def normalize_cell_range(value: str, max_row: int) -> str:
    cleaned = clean_cell_ref(value)
    # Some SpreadsheetBench labels use BD2:308 to mean BD2:BD308.
    cleaned = re.sub(r"^([A-Za-z]+)(\d+):(\d+)$", r"\1\2:\1\3", cleaned)
    min_col, min_row, max_col, max_row_bound = range_boundaries(cleaned)
    if min_row is None:
        min_row = 1
    if max_row_bound is None:
        max_row_bound = max_row
    return f"{min_col},{min_row},{max_col},{max_row_bound}"


def parse_sheet_list(value: str) -> list[str]:
    raw = clean_sheet_name(value)
    if not raw:
        return []
    return [clean_sheet_name(part) for part in raw.split(",") if clean_sheet_name(part)]


def split_sheet_range(sheet_name: str, cell_range: str) -> tuple[str, str]:
    rng = str(cell_range or "").strip()
    if "!" in rng:
        left, right = rng.split("!", 1)
        return clean_sheet_name(left) or sheet_name, clean_cell_ref(right)
    return sheet_name, clean_cell_ref(rng)


def split_reference_segments(default_sheet: str, cell_range: str) -> list[tuple[str, str]]:
    """SpreadsheetBench sometimes stores comma-separated refs across sheets."""
    segments: list[tuple[str, str]] = []
    current_sheet = default_sheet
    for raw_part in [item.strip() for item in str(cell_range or "").split(",") if item.strip()]:
        current_sheet, current_range = split_sheet_range(current_sheet, raw_part)
        segments.append((current_sheet, current_range))
    return segments


def resolve_workbook_sheet(sheet_name: str, sheetnames: list[str]) -> str:
    if sheet_name in sheetnames:
        return sheet_name
    stripped = sheet_name.strip()
    for actual in sheetnames:
        if actual.strip() == stripped:
            return actual
    folded = " ".join(stripped.lower().replace("，", ",").split())
    for actual in sheetnames:
        actual_folded = " ".join(actual.strip().lower().replace("，", ",").split())
        if actual_folded == folded:
            return actual
        if actual_folded.startswith(folded + ","):
            return actual
    raise RuntimeError(f"answer sheet not found: {sheet_name}")


def load_candidate_object(candidate_path: str) -> dict:
    raw = load_json_file(candidate_path)
    if not isinstance(raw, dict):
        raise RuntimeError("candidate json must be an object")
    return raw


def resolve_expected_location(
    answer_path: str,
    sheet_name: str | None,
    cell_range: str | None,
    candidate: dict,
) -> tuple[str, str]:
    candidate_sheet = candidate.get("answer_sheet")
    candidate_range = candidate.get("answer_position")
    resolved_range = str(cell_range or candidate_range or "").strip()
    resolved_sheet = clean_sheet_name(sheet_name)
    if resolved_range:
        resolved_sheet, resolved_range = split_sheet_range(resolved_sheet, resolved_range)
    if not resolved_sheet and isinstance(candidate_sheet, str):
        resolved_sheet = clean_sheet_name(candidate_sheet)
    if not resolved_range:
        raise RuntimeError("answer range is missing; provide --range or candidate answer_position")
    wb = load_workbook(answer_path, data_only=True, read_only=True)
    try:
        sheetnames = list(wb.sheetnames)
    finally:
        wb.close()
    if not resolved_sheet:
        if len(sheetnames) == 1:
            resolved_sheet = sheetnames[0]
        else:
            raise RuntimeError("answer sheet is missing; provide --sheet or candidate answer_sheet")
    return resolved_sheet, resolved_range


def read_expected(answer_path: str, sheet_name: str, cell_range: str) -> list[list[Any]]:
    wb = load_workbook(answer_path, data_only=True, read_only=True)
    rows: list[list[Any]] = []
    for segment_sheet, segment_range in split_reference_segments(sheet_name, cell_range):
        target_sheets: list[str]
        try:
            target_sheets = [resolve_workbook_sheet(segment_sheet, list(wb.sheetnames))]
        except RuntimeError:
            sheet_list = parse_sheet_list(segment_sheet)
            if len(sheet_list) <= 1:
                raise
            target_sheets = [resolve_workbook_sheet(item, list(wb.sheetnames)) for item in sheet_list]
        for actual_sheet in target_sheets:
            ws = wb[actual_sheet]
            min_col, min_row, max_col, max_row = [int(v) for v in normalize_cell_range(segment_range, ws.max_row).split(",")]
            for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
                rows.append([norm(cell.value) for cell in row])
    wb.close()
    return rows


def candidate_values_payload(candidate: dict) -> Any:
    for key in ["values", "matrix", "data", "answer", "value", "result"]:
        if key in candidate:
            return candidate.get(key)
    raise RuntimeError("candidate json must contain values/matrix/data/answer/value/result")


def read_candidate(candidate: dict, expected: list[list[Any]]) -> list[list[Any]]:
    values = candidate_values_payload(candidate)
    expected_rows = len(expected)
    expected_cols = len(expected[0]) if expected and expected[0] else 0
    if not isinstance(values, list):
        if expected_rows == 1 and expected_cols == 1:
            return [[norm(values)]]
        raise RuntimeError("candidate scalar value is only accepted for a 1x1 expected range")
    if values and all(not isinstance(row, list) for row in values):
        if expected_rows == 1:
            return [[norm(v) for v in values]]
        if expected_cols == 1 and len(values) == expected_rows:
            return [[norm(v)] for v in values]
        raise RuntimeError("candidate flat list is ambiguous; use values: list[list]")
    out: list[list[Any]] = []
    for row in values:
        if not isinstance(row, list):
            raise RuntimeError("candidate values must be list[list]")
        out.append([norm(v) for v in row])
    return out


def input_workbook_from_payload(payload: dict) -> str | None:
    for item in payload.get("materials", []):
        if not isinstance(item, dict):
            continue
        value = item.get("path")
        role = item.get("role")
        if (
            isinstance(value, str)
            and role == "source-artifact"
            and value.lower().endswith((".xlsx", ".xlsm"))
            and os.path.exists(value)
        ):
            return value
    for item in payload.get("materials", []):
        if not isinstance(item, dict):
            continue
        value = item.get("path")
        if isinstance(value, str) and value.lower().endswith((".xlsx", ".xlsm")) and os.path.exists(value):
            return value
    return None


def simple_cell_refs(formula: Any) -> list[str]:
    if not isinstance(formula, str) or not formula.startswith("="):
        return []
    refs: list[str] = []
    # Deliberately narrow: this is a consistency guard for same-sheet formulas,
    # not a full Excel parser.
    for match in re.finditer(r"(?<![A-Za-z0-9_])\$?([A-Za-z]{1,3})\$?(\d+)", formula):
        ref = f"{match.group(1).upper()}{match.group(2)}"
        if ref not in refs:
            refs.append(ref)
    return refs


def detect_dataset_formula_selector_mismatch(
    payload: dict,
    answer_path: str,
    sheet_name: str,
    cell_range: str,
) -> list[str]:
    input_path = input_workbook_from_payload(payload)
    if not input_path:
        return []
    try:
        input_wb = load_workbook(input_path, data_only=False, read_only=False)
        answer_wb = load_workbook(answer_path, data_only=False, read_only=False)
    except Exception:
        return []
    try:
        target_refs = set()
        mismatches: list[str] = []
        for segment_sheet, segment_range in split_reference_segments(sheet_name, cell_range):
            input_sheet = resolve_workbook_sheet(segment_sheet, list(input_wb.sheetnames))
            answer_sheet = resolve_workbook_sheet(segment_sheet, list(answer_wb.sheetnames))
            input_ws = input_wb[input_sheet]
            answer_ws = answer_wb[answer_sheet]
            min_col, min_row, max_col, max_row = [int(v) for v in normalize_cell_range(segment_range, input_ws.max_row).split(",")]
            for row in input_ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
                for cell in row:
                    for ref in simple_cell_refs(cell.value):
                        target_refs.add((input_sheet, answer_sheet, ref))
        for input_sheet, answer_sheet, ref in sorted(target_refs):
            try:
                input_ws = input_wb[input_sheet]
                answer_ws = answer_wb[answer_sheet]
                input_value = norm(input_ws[ref].value)
                answer_value = norm(answer_ws[ref].value)
            except Exception:
                continue
            if input_value != answer_value:
                mismatches.append(f"formula reference {input_sheet}!{ref} differs between public input and hidden answer: input={input_value!r} answer={answer_value!r}")
                if len(mismatches) >= 8:
                    break
        return mismatches
    finally:
        input_wb.close()
        answer_wb.close()


def compare_rows(expected: list[list[Any]], got: list[list[Any]]) -> tuple[bool, list[str]]:
    details: list[str] = []
    if len(expected) != len(got):
        details.append(f"row count mismatch: expected={len(expected)} got={len(got)}")
        return False, details
    for r, (erow, grow) in enumerate(zip(expected, got), start=1):
        if len(erow) != len(grow):
            details.append(f"col count mismatch at row {r}: expected={len(erow)} got={len(grow)}")
            return False, details
        for c, (ev, gv) in enumerate(zip(erow, grow), start=1):
            if ev != gv:
                details.append(f"mismatch at r{r}c{c}: expected={ev!r} got={gv!r}")
                if len(details) >= 20:
                    return False, details
    return len(details) == 0, details


def main() -> int:
    args = parse_args()
    try:
        payload = load_payload(args.input_json)
        candidate_path = resolve_path(args.candidate, payload, prefer_payload_output=True)
        answer_path = resolve_path(args.answer, payload)
        if not os.path.exists(candidate_path):
            result = {
                "passed": False,
                "blocking": True,
                "summary": "Candidate file not found.",
                "details": [f"missing candidate: {candidate_path}"],
            }
            print(json.dumps(result, ensure_ascii=False))
            return 0
        candidate = load_candidate_object(candidate_path)
        sheet_name, cell_range = resolve_expected_location(answer_path, args.sheet, args.cell_range, candidate)
        expected = read_expected(answer_path, sheet_name, cell_range)
        got = read_candidate(candidate, expected)
        passed, diffs = compare_rows(expected, got)
        dataset_mismatches = [] if passed else detect_dataset_formula_selector_mismatch(payload, answer_path, sheet_name, cell_range)
        result = {
            "passed": passed,
            "blocking": True,
            "failureKind": "dataset_failure" if dataset_mismatches else None,
            "summary": "Spreadsheet answer matched." if passed else ("Spreadsheet dataset consistency failure." if dataset_mismatches else "Spreadsheet answer mismatch."),
            "details": (dataset_mismatches + diffs) if dataset_mismatches else (diffs if diffs else [f"matched sheet={sheet_name} range={cell_range}"]),
        }
        print(json.dumps(result, ensure_ascii=False))
        return 0
    except Exception as e:
        result = {
            "passed": False,
            "blocking": True,
            "summary": "Validator crashed.",
            "details": [str(e)],
        }
        print(json.dumps(result, ensure_ascii=False))
        return 0


if __name__ == "__main__":
    sys.exit(main())
