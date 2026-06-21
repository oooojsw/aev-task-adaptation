from __future__ import annotations

import argparse
import importlib.util
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook


EXPECTED_EXPERIMENTS = {
    "officeqa172-a3b-with-skill": (172, 93),
    "officeqa172-a3b-no-skill": (172, 81),
    "officeqa172-qwen35-4b-with-skill": (172, 13),
    "officeqa172-qwen35-4b-no-skill": (172, 14),
}

REQUIRED_SHEETS = {
    "Dashboard",
    "Official Experiments",
    "Run Ledger",
    "Case Results",
    "Official Case Results",
    "Errors",
    "Fields",
}


def rows_as_dicts(sheet):
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value or "") for value in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:]]


def main() -> None:
    auto_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description="Smoke-test the Direct Chat master report workbook.")
    parser.add_argument(
        "--report",
        type=Path,
        default=auto_root / "runtime-state" / "direct-chat-reports" / "direct-chat-evaluation-master.xlsx",
    )
    args = parser.parse_args()

    workbook = load_workbook(args.report, read_only=False, data_only=False)
    missing = REQUIRED_SHEETS.difference(workbook.sheetnames)
    assert not missing, f"Missing sheets: {sorted(missing)}"

    official_rows = rows_as_dicts(workbook["Official Experiments"])
    official = {str(row.get("experimentID", "")): row for row in official_rows}
    for experiment_id, (expected_total, expected_passed) in EXPECTED_EXPERIMENTS.items():
        assert experiment_id in official, f"Missing official experiment: {experiment_id}"
        row = official[experiment_id]
        assert int(row["total"]) == expected_total, (experiment_id, row["total"])
        assert int(row["passed"]) == expected_passed, (experiment_id, row["passed"])
        assert not bool(row["isComplete"]), f"Legacy truncated experiment must not remain complete: {experiment_id}"
        assert row.get("validityStatus") == "invalid", f"Legacy experiment must be marked invalid: {experiment_id}"

    run_rows = rows_as_dicts(workbook["Run Ledger"])
    run_keys = [(row.get("runID"), row.get("modelID")) for row in run_rows]
    assert len(run_keys) == len(set(run_keys)), "Run Ledger contains duplicate run/model rows"

    case_rows = rows_as_dicts(workbook["Case Results"])
    case_keys = [(row.get("runID"), row.get("modelID"), row.get("caseID")) for row in case_rows]
    assert len(case_keys) == len(set(case_keys)), "Case Results contains duplicate run/model/case rows"

    assert workbook["Dashboard"]._charts, "Dashboard chart is missing"

    with tempfile.TemporaryDirectory(prefix="direct-chat-master-report-smoke-") as temp_value:
        temp_root = Path(temp_value)
        runs_root = temp_root / "runs"
        good_run = runs_root / "good-run"
        retry_run = runs_root / "retry-run"
        broken_run = runs_root / "broken-run"
        good_run.mkdir(parents=True)
        retry_run.mkdir(parents=True)
        broken_run.mkdir(parents=True)

        (good_run / "summary.json").write_text(
            json.dumps(
                {
                    "schema": "direct-chat-summary/v1",
                    "validity": {"status": "valid", "reasons": []},
                    "summary": [
                        {
                            "modelID": "mock-model__no-skill",
                            "total": 1,
                            "passed": 1,
                            "failed": 0,
                            "passRate": 1,
                        }
                    ],
                    "results": [
                        {
                            "caseID": "CASE001",
                            "modelID": "mock-model__no-skill",
                            "provider": "mock",
                            "skillMode": "no-skill",
                            "passed": True,
                            "score": 1,
                            "answer": "=HYPERLINK(\"https://invalid.example\")\u0001",
                            "errorCategory": "none",
                            "attempts": 1,
                            "latencyMs": 10,
                        }
                    ],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        (good_run / "config.snapshot.json").write_text("{broken config", encoding="utf-8")
        (retry_run / "summary.json").write_text(
            json.dumps(
                {
                    "schema": "direct-chat-summary/v1",
                    "validity": {"status": "valid", "reasons": []},
                    "summary": [
                        {
                            "modelID": "mock-model__no-skill",
                            "total": 1,
                            "passed": 0,
                            "failed": 1,
                            "passRate": 0,
                        }
                    ],
                    "results": [
                        {
                            "caseID": "CASE001",
                            "modelID": "mock-model__no-skill",
                            "provider": "mock",
                            "skillMode": "no-skill",
                            "passed": False,
                            "score": 0,
                            "answer": "",
                            "errorCategory": "auth",
                            "error": "simulated quota failure",
                        }
                    ],
                }
            ),
            encoding="utf-8",
        )
        (broken_run / "summary.json").write_text("{not json", encoding="utf-8")
        composites_path = temp_root / "composites.json"
        composites_path.write_text(
            json.dumps(
                {
                    "experiments": [
                        {
                            "id": "robustness",
                            "label": "Robustness",
                            "expectedCases": 1,
                            "replacePolicy": "prefer-non-infrastructure",
                            "sources": ["good-run", "retry-run"],
                        }
                    ]
                }
            ),
            encoding="utf-8",
        )
        generated_report = temp_root / "master.xlsx"
        subprocess.run(
            [
                sys.executable,
                str(auto_root / "scripts" / "update-direct-chat-master-report.py"),
                "--runs-root",
                str(runs_root),
                "--out",
                str(generated_report),
                "--composites-config",
                str(composites_path),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
        injected = load_workbook(generated_report, read_only=False, data_only=False)
        injected_runs = rows_as_dicts(injected["Run Ledger"])
        assert any(row.get("status") == "summary-read-error" for row in injected_runs)
        assert any(row.get("reportWarnings") for row in injected_runs if row.get("runID") == "good-run")
        injected_cases = rows_as_dicts(injected["Case Results"])
        safe_answer = next(row["answer"] for row in injected_cases if row.get("caseID") == "CASE001" and row.get("passed"))
        assert str(safe_answer).startswith("'="), "Formula-like answer was not neutralized"
        assert "\u0001" not in str(safe_answer), "Illegal XML control character was not removed"
        injected_official = rows_as_dicts(injected["Official Experiments"])
        assert injected_official[0]["passed"] == 1, "Infrastructure retry replaced a valid result"
        assert injected_official[0]["isComplete"] is True
        status = json.loads((temp_root / "master-report-status.json").read_text(encoding="utf-8"))
        assert status["status"] == "updated"

        module_path = auto_root / "scripts" / "update-direct-chat-master-report.py"
        spec = importlib.util.spec_from_file_location("direct_chat_master_report_smoke_module", module_path)
        assert spec and spec.loader
        report_module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(report_module)
        locked_out = temp_root / "locked-master.xlsx"
        locked_out.write_bytes(b"existing")
        original_replace = report_module.os.replace
        original_sleep = report_module.time.sleep

        def simulated_replace(source, destination):
            if Path(destination) == locked_out:
                raise PermissionError("simulated Excel lock")
            return original_replace(source, destination)

        report_module.os.replace = simulated_replace
        report_module.time.sleep = lambda _seconds: None
        pending_workbook = Workbook()
        pending_workbook.active["A1"] = "pending"
        pending_status = report_module.save_atomic(pending_workbook, locked_out)
        assert pending_status == "pending-excel-close"
        assert (temp_root / "locked-master.pending.xlsx").exists()
        report_module.os.replace = original_replace
        report_module.time.sleep = original_sleep
        assert report_module.promote_pending(locked_out)
        promoted = load_workbook(locked_out)
        assert promoted.active["A1"].value == "pending"

    print(
        f"master-report-smoke-ok sheets={len(workbook.sheetnames)} "
        f"runs={len(run_rows)} cases={len(case_rows)} official={len(official_rows)}"
    )


if __name__ == "__main__":
    main()
